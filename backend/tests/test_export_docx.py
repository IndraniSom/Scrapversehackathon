"""Professional DOCX export: deterministic order, styles, unicode, overflow."""

import hashlib
import io

from docx import Document

from backend.export_docx import build_proposal_docx


def _sample(sections):
    """Build minimal proposal dict with given sections."""
    return {
        "proposal_id": "prop-001",
        "title": "Pond Monitoring \u20b9 90L — \u0939\u093f\u0902\u0926\u0940",
        "revision": 2,
        "sections": sections,
    }


def test_docx_deterministic_order():
    """Shuffled sections produce identical bytes."""
    secs_a = [
        {"title": "Approach", "body": "Body A", "citation": "Doc 1 p1", "order": 1, "state": "APPROVED"},
        {"title": "Cover", "body": "Body B", "citation": "Doc 1 p2", "order": 0, "state": "APPROVED"},
    ]
    secs_b = list(reversed(secs_a))
    a = build_proposal_docx(_sample(secs_a))
    b = build_proposal_docx(_sample(secs_b))
    assert hashlib.sha256(a).hexdigest() == hashlib.sha256(b).hexdigest()


def test_docx_uses_styles_not_direct_formatting():
    """Paragraphs use named styles, not direct bold/italic formatting."""
    data = _sample([
        {"title": "Section A", "body": "Content", "citation": "Doc p1", "order": 0, "state": "APPROVED"},
    ])
    doc = Document(io.BytesIO(build_proposal_docx(data)))
    styles = {p.style.name for p in doc.paragraphs if p.text.strip()}
    assert "Heading 1" in styles
    assert "Normal" in styles
    assert "Title" in styles
    for para in doc.paragraphs:
        for run in para.runs:
            # direct bold/italic should not be set (None means inherit from style)
            assert run.bold is None or run.bold is False
            # font size should be inherited
            assert run.font.size is None


def test_docx_unicode_inr_preserved():
    """INR symbol and Hindi text round-trip through DOCX."""
    data = _sample([
        {"title": "Financial \u20b9", "body": "Amount \u20b9 12,00,000 and \u0939\u093f\u0902\u0926\u0940 clause", "citation": "Doc p1 \u20b9", "order": 0, "state": "APPROVED"},
    ])
    doc = Document(io.BytesIO(build_proposal_docx(data)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "\u20b9" in text
    assert "\u0939\u093f\u0902\u0926\u0940" in text


def test_docx_table_overflow_handled():
    """Very long body is truncated without crashing and remains inspectable."""
    long_body = "A" * 8000 + "\n" + "B" * 5000
    data = _sample([
        {"title": "Overflow", "body": long_body, "citation": "Doc p99", "order": 0, "state": "APPROVED"},
    ])
    raw = build_proposal_docx(data)
    doc = Document(io.BytesIO(raw))
    assert len(doc.paragraphs) > 5
    # contains truncation marker
    assert any("truncated" in p.text for p in doc.paragraphs)


def test_docx_headers_footers_toc_citations():
    """Header, footer, TOC, citations and revision metadata are present."""
    data = _sample([
        {"title": "Alpha", "body": "Alpha body", "citation": "Tender p.5", "order": 0, "state": "APPROVED"},
        {"title": "Beta", "body": "Beta body", "citation": "Tender p.6", "order": 1, "state": "APPROVED"},
    ])
    raw = build_proposal_docx(data)
    doc = Document(io.BytesIO(raw))
    # header/footer
    assert doc.sections[0].header.paragraphs[0].text.startswith("Pond Monitoring")
    assert "Page" in doc.sections[0].footer.paragraphs[0].text
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Contents" in full
    assert "Source: Tender p.5" in full
    assert "Source: Tender p.6" in full
    assert "Rev 2" in full
    assert "State: APPROVED" in full


def test_docx_visual_inspect_valid():
    """Generated DOCX is valid zip and contains expected structure."""
    data = _sample([
        {"title": "Visual", "body": "Inspect me", "order": 0, "state": "APPROVED"},
    ])
    raw = build_proposal_docx(data)
    assert raw[:2] == b"PK"  # zip signature
    doc = Document(io.BytesIO(raw))
    assert doc.core_properties is not None
